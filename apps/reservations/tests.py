import secrets
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.guests.models import Guest, normalize_phone
from apps.reservations import services
from apps.reservations.models import Reservation, TableOccupancy, WaitlistEntry
from apps.venues.models import Area, DurationRule, Shift, Table, TableCombination, Venue

LIMA = ZoneInfo("America/Lima")
STARTS_AT = "starts_at"


class BaseSalaTestCase(TestCase):
    """Sala minima: un salon con tres mesas y un turno de cena."""

    @classmethod
    def setUpTestData(cls):
        cls.venue = Venue.objects.create(name="Zisa", timezone="America/Lima")
        cls.area = Area.objects.create(venue=cls.venue, name="Salon")
        cls.t1 = Table.objects.create(area=cls.area, code="S1", min_seats=2, max_seats=4)
        cls.t2 = Table.objects.create(area=cls.area, code="S2", min_seats=2, max_seats=4)
        cls.t3 = Table.objects.create(area=cls.area, code="S3", min_seats=6, max_seats=8)

        cls.combo = TableCombination.objects.create(
            venue=cls.venue, name="S1+S2", min_seats=5, max_seats=8
        )
        cls.combo.tables.set([cls.t1, cls.t2])

        Shift.objects.create(
            venue=cls.venue, name="Cena", weekdays="0,1,2,3,4,5,6",
            start_time=time(19, 0), last_seating=time(21, 45), end_time=time(23, 0),
        )
        DurationRule.objects.create(venue=cls.venue, min_party=1, max_party=4, minutes=90)
        DurationRule.objects.create(venue=cls.venue, min_party=5, max_party=8, minutes=120)

        cls.guest = Guest.objects.create(phone="+51987654321", name="Ana Rojas")

    def cena(self, dias_desde_hoy=1, hora=20, minuto=0):
        """Un inicio valido dentro del turno de cena, en hora de Lima."""
        dia = timezone.localtime(timezone.now(), LIMA).date() + timedelta(
            days=dias_desde_hoy
        )
        return datetime.combine(dia, time(hora, minuto), tzinfo=LIMA)


class TelefonoTests(TestCase):
    def test_normaliza_las_formas_que_escribe_la_gente(self):
        esperado = "+51987654321"
        for entrada in ["987654321", "987 654 321", "+51 987-654-321",
                        "0051987654321", "(51) 987654321"]:
            with self.subTest(entrada=entrada):
                self.assertEqual(normalize_phone(entrada), esperado)

    def test_el_cliente_se_guarda_normalizado(self):
        guest = Guest.objects.create(phone="987 654 999", name="Luis Salas")
        self.assertEqual(guest.phone, "+51987654999")


class DisponibilidadTests(BaseSalaTestCase):
    def test_elige_la_mesa_con_menos_desperdicio(self):
        inicio = self.cena()
        ofertas = services.find_availability(self.venue, inicio,
                                             inicio + timedelta(minutes=90), 2)
        # S3 es de 6-8, no acepta un grupo de 2; S1 y S2 si.
        self.assertTrue(ofertas)
        self.assertEqual(ofertas[0].slack, 2)
        self.assertIn(ofertas[0].codes, {"S1", "S2"})

    def test_usa_una_combinacion_cuando_ninguna_mesa_alcanza_sola(self):
        inicio = self.cena()
        # Ocupamos S3, la unica mesa que sola aguanta 6.
        services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=inicio, party_size=6,
            source=Reservation.Source.PHONE, tables=[self.t3],
        )
        otro = Guest.objects.create(phone="+51900000001", name="Bruno Vega")
        reserva = services.create_reservation(
            venue=self.venue, guest=otro, starts_at=inicio, party_size=6,
            source=Reservation.Source.PHONE,
        )
        self.assertEqual(reserva.table_codes, "S1, S2")

    def test_no_hay_mesa_cuando_todo_esta_ocupado(self):
        inicio = self.cena()
        for indice, mesa in enumerate([self.t1, self.t2, self.t3]):
            guest = Guest.objects.create(phone=f"+5190000010{indice}", name=f"G{indice}")
            services.create_reservation(
                venue=self.venue, guest=guest, starts_at=inicio, party_size=2,
                source=Reservation.Source.PHONE, tables=[mesa],
            )
        otro = Guest.objects.create(phone="+51900000199", name="Sin suerte")
        with self.assertRaises(services.NoAvailability):
            services.create_reservation(
                venue=self.venue, guest=otro, starts_at=inicio, party_size=2,
                source=Reservation.Source.PHONE,
            )


class SolapamientoTests(BaseSalaTestCase):
    def test_la_base_rechaza_dos_ocupaciones_activas_que_se_pisan(self):
        """La rechaza Postgres, no la capa de servicio."""
        inicio = self.cena()
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=inicio, party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                TableOccupancy.objects.create(
                    table=self.t1, reservation=reserva,
                    starts_at=inicio + timedelta(minutes=30),
                    ends_at=inicio + timedelta(minutes=120),
                )

    def test_cancelar_libera_la_mesa_de_inmediato(self):
        inicio = self.cena()
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=inicio, party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        services.cancel(reserva, reason="El cliente llamo", user=None)

        otro = Guest.objects.create(phone="+51900000002", name="Marta Quispe")
        nueva = services.create_reservation(
            venue=self.venue, guest=otro, starts_at=inicio, party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        self.assertEqual(nueva.table_codes, "S1")

    def test_un_bloqueo_impide_reservar_encima(self):
        inicio = self.cena()
        services.block_table(self.t1, inicio, inicio + timedelta(hours=3),
                             reason="Toldo roto")
        ofertas = services.find_availability(self.venue, inicio,
                                             inicio + timedelta(minutes=90), 2)
        self.assertNotIn("S1", [o.codes for o in ofertas])

    def test_reservas_consecutivas_no_se_pisan(self):
        """Un rango termina donde empieza el siguiente: tstzrange es [inicio, fin)."""
        inicio = self.cena(hora=19)
        services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=inicio, party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        otro = Guest.objects.create(phone="+51900000003", name="Rita Ibanez")
        segunda = services.create_reservation(
            venue=self.venue, guest=otro, starts_at=inicio + timedelta(minutes=90),
            party_size=2, source=Reservation.Source.PHONE, tables=[self.t1],
        )
        self.assertEqual(segunda.table_codes, "S1")


class EstadoTests(BaseSalaTestCase):
    def _reserva(self):
        return services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )

    def test_recorrido_completo_deja_rastro(self):
        reserva = self._reserva()
        services.seat(reserva)
        reserva = services.finish(reserva)

        self.assertEqual(reserva.status, Reservation.Status.FINISHED)
        self.assertIsNotNone(reserva.seated_at)
        self.assertIsNotNone(reserva.released_at)
        self.assertEqual(reserva.events.count(), 3)  # creada, sentada, finalizada
        self.assertFalse(reserva.occupancies.filter(is_active=True).exists())

    def test_no_se_puede_sentar_una_reserva_cancelada(self):
        reserva = self._reserva()
        services.cancel(reserva, reason="prueba")
        with self.assertRaises(services.InvalidTransition):
            services.seat(reserva)

    def test_el_no_show_suma_al_contador_del_cliente(self):
        reserva = self._reserva()
        services.mark_no_show(reserva)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.no_shows, 1)
        self.assertTrue(self.guest.is_risky is False)  # hace falta un segundo planton


class CalendarioTests(BaseSalaTestCase):
    def test_rechaza_una_hora_fuera_de_turno(self):
        fuera = self.cena(hora=17)
        with self.assertRaises(services.NoAvailability):
            services.create_reservation(
                venue=self.venue, guest=self.guest, starts_at=fuera, party_size=2,
                source=Reservation.Source.PHONE,
            )

    def test_rechaza_despues_de_la_ultima_entrada(self):
        tarde = self.cena(hora=22, minuto=30)
        with self.assertRaises(services.NoAvailability):
            services.create_reservation(
                venue=self.venue, guest=self.guest, starts_at=tarde, party_size=2,
                source=Reservation.Source.PHONE,
            )

    def test_la_madrugada_pertenece_al_servicio_de_la_noche_anterior(self):
        medianoche = datetime.combine(date(2026, 9, 12), time(0, 30), tzinfo=LIMA)
        self.assertEqual(
            services.service_date_for(self.venue, medianoche), date(2026, 9, 11)
        )

    def test_el_grupo_grande_no_pasa_por_web(self):
        with self.assertRaises(services.ReservationError):
            services.create_reservation(
                venue=self.venue, guest=self.guest, starts_at=self.cena(),
                party_size=20, source=Reservation.Source.WEB,
            )


class BorradoTests(BaseSalaTestCase):
    """Borrar una reserva no debe exigir permiso sobre "movimiento", que no se da a nadie."""

    def setUp(self):
        from django.contrib.auth.models import Permission, User

        # Clave aleatoria en cada ejecucion: el test entra con force_login.
        self.staff = User.objects.create_user("host", password=secrets.token_urlsafe(),
                                              is_staff=True)
        self.staff.user_permissions.add(
            Permission.objects.get(codename="view_reservation"),
            Permission.objects.get(codename="delete_reservation"),
        )
        self.reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.RECEPCION,
        )
        services.seat(self.reserva)  # nace confirmada; sentar deja otro movimiento

    def test_staff_con_permiso_puede_borrar_la_reserva_y_su_historial(self):
        from django.urls import reverse

        from apps.reservations.models import ReservationEvent

        self.client.force_login(self.staff)
        url = reverse("admin:reservations_reservation_delete", args=[self.reserva.pk])

        pagina = self.client.get(url)
        self.assertEqual(pagina.status_code, 200)
        self.assertFalse(pagina.context["perms_lacking"])

        self.client.post(url, {"post": "yes"})
        self.assertFalse(Reservation.objects.filter(pk=self.reserva.pk).exists())
        self.assertFalse(ReservationEvent.objects.filter(reservation_id=self.reserva.pk).exists())
        self.assertFalse(TableOccupancy.objects.filter(reservation_id=self.reserva.pk).exists())

    def test_los_movimientos_siguen_sin_poder_borrarse_sueltos(self):
        from django.contrib import admin

        from apps.reservations.models import ReservationEvent

        evento = ReservationEvent.objects.filter(reservation=self.reserva).first()
        ma = admin.site._registry[ReservationEvent]
        self.assertFalse(ma.has_delete_permission(None, evento))


class ReservaWebTests(BaseSalaTestCase):
    """Los campos llevan los mismos nombres que en el Contact Form 7 de zisa.pe."""

    TOKEN = "clave-de-prueba"

    def post(self, token=TOKEN, **campos):
        from django.test import override_settings

        datos = {
            "full-name": "Lucia Paredes",
            "your-phone": "987 111 222",
            "num-person": "2",
            "date-reservation": self.cena().date().isoformat(),
            "time-field": "08:00 PM",
            "indications": "Mesa junto a la ventana",
        }
        datos.update({k.replace("_", "-"): v for k, v in campos.items()})
        cabeceras = {"HTTP_AUTHORIZATION": f"Bearer {token}"} if token else {}
        with override_settings(WEB_RESERVATION_TOKEN=self.TOKEN):
            return self.client.post("/api/reservas/web/", datos, **cabeceras)

    def test_crea_la_reserva_web_con_mesa_y_cliente(self):
        r = self.post()
        self.assertEqual(r.status_code, 201, r.content)
        cuerpo = r.json()
        self.assertTrue(cuerpo["ok"])

        reserva = Reservation.objects.get(code=cuerpo["code"])
        self.assertEqual(reserva.source, Reservation.Source.WEB)
        self.assertEqual(reserva.party_size, 2)
        self.assertEqual(reserva.guest.phone, "+51987111222")
        self.assertEqual(reserva.guest_notes, "Mesa junto a la ventana")
        # "08:00 PM" del desplegable son las 20:00 de Lima.
        self.assertEqual(timezone.localtime(reserva.starts_at, LIMA).hour, 20)
        self.assertNotEqual(reserva.table_codes, "sin asignar")
        self.assertIn(cuerpo["code"], cuerpo["message"])

    def test_cuerpo_vacio_usa_los_campos_de_la_url(self):
        # Asi llega desde el hosting de zisa.pe: cuerpo vacio, datos en la URL.
        from urllib.parse import urlencode

        from django.test import override_settings

        datos = {
            "full-name": "Lucia Paredes", "your-phone": "987 111 222",
            "num-person": "2", "date-reservation": self.cena().date().isoformat(),
            "time-field": "08:00 PM", "indications": "",
        }
        with override_settings(WEB_RESERVATION_TOKEN=self.TOKEN):
            r = self.client.generic(
                "POST", f"/api/reservas/web/?{urlencode(datos)}", b"",
                content_type="text/plain", HTTP_AUTHORIZATION=f"Bearer {self.TOKEN}",
            )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Reservation.objects.get().guest.name, "Lucia Paredes")

    def test_sin_clave_o_con_otra_no_entra(self):
        self.assertEqual(self.post(token=None).status_code, 401)
        self.assertEqual(self.post(token="otra").status_code, 401)
        self.assertFalse(Reservation.objects.exists())

    def test_sin_clave_configurada_la_entrada_esta_apagada(self):
        from django.test import override_settings

        with override_settings(WEB_RESERVATION_TOKEN=""):
            r = self.client.post("/api/reservas/web/", {},
                                 HTTP_AUTHORIZATION="Bearer ")
        self.assertEqual(r.status_code, 503)

    def test_datos_incompletos_dicen_que_campo_falla(self):
        r = self.post(time_field="Hora de reserva", full_name="")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(set(r.json()["errors"]), {"time-field", "full-name"})

    def test_acepta_hora_en_24h(self):
        r = self.post(time_field="20:30")
        self.assertEqual(r.status_code, 201, r.content)

    def test_sin_mesa_devuelve_el_motivo(self):
        # S3 pide minimo 6: la tercera pareja a la misma hora ya no cabe.
        for telefono in ("987000001", "987000002"):
            self.assertEqual(self.post(your_phone=telefono).status_code, 201)
        r = self.post(your_phone="987000003")
        self.assertEqual(r.status_code, 409)
        self.assertIn("No hay mesa", r.json()["message"])

    def test_fuera_de_turno_devuelve_el_motivo(self):
        r = self.post(time_field="07:00 AM")
        self.assertEqual(r.status_code, 409)
        self.assertIn("turno", r.json()["message"])

    def test_reenviar_el_formulario_no_duplica(self):
        primera = self.post().json()
        r = self.post()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["code"], primera["code"])
        self.assertEqual(Reservation.objects.count(), 1)

    def test_un_rechazo_no_deja_cliente_creado(self):
        r = self.post(num_person="30", your_phone="987 333 111")
        self.assertEqual(r.status_code, 409)
        self.assertFalse(Guest.objects.filter(phone="+51987333111").exists())

    def test_grupo_grande_se_deriva_al_telefono(self):
        r = self.post(num_person="30")
        self.assertEqual(r.status_code, 409)
        self.assertIn("teléfono", r.json()["message"])


class TareaNoShowTests(BaseSalaTestCase):
    TOKEN = "clave-cron"

    def llamar(self, token=TOKEN):
        from django.test import override_settings

        cabeceras = {"HTTP_AUTHORIZATION": f"Bearer {token}"} if token else {}
        with override_settings(CRON_TOKEN=self.TOKEN):
            return self.client.post("/api/tareas/no-shows/", **cabeceras)

    def reserva_atrasada(self, minutos):
        """Reserva confirmada que empezo hace `minutos` y nadie llego."""
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.RECEPCION,
        )
        self.assertEqual(reserva.status, Reservation.Status.CONFIRMED)
        inicio = timezone.now() - timedelta(minutes=minutos)
        Reservation.objects.filter(pk=reserva.pk).update(
            starts_at=inicio, ends_at=inicio + timedelta(minutes=90))
        return reserva

    def test_marca_solo_las_que_pasaron_la_gracia(self):
        vencida = self.reserva_atrasada(minutos=20)
        a_tiempo = self.reserva_atrasada(minutos=5)

        r = self.llamar()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["total"], 1)

        vencida.refresh_from_db()
        a_tiempo.refresh_from_db()
        self.assertEqual(vencida.status, Reservation.Status.NO_SHOW)
        self.assertEqual(a_tiempo.status, Reservation.Status.CONFIRMED)

    def test_sin_clave_o_con_otra_no_entra(self):
        self.reserva_atrasada(minutos=20)
        self.assertEqual(self.llamar(token=None).status_code, 401)
        self.assertEqual(self.llamar(token="otra").status_code, 401)
        self.assertFalse(Reservation.objects.filter(
            status=Reservation.Status.NO_SHOW).exists())

    def test_sin_clave_configurada_la_ruta_esta_apagada(self):
        from django.test import override_settings

        with override_settings(CRON_TOKEN=""):
            r = self.client.post("/api/tareas/no-shows/", HTTP_AUTHORIZATION="Bearer ")
        self.assertEqual(r.status_code, 503)


class WalkInTests(BaseSalaTestCase):
    def ahora(self):
        return self.cena(hora=20)  # un "ahora" fijo dentro del turno de cena

    def test_sienta_en_el_acto_y_ocupa_la_mesa(self):
        r = services.seat_walk_in(self.venue, self.t1, 3, now=self.ahora())
        self.assertEqual(r.status, Reservation.Status.SEATED)
        self.assertEqual(r.source, Reservation.Source.RECEPCION)
        self.assertEqual(r.guest.name, services.WALK_IN_NAME)
        self.assertEqual(r.table_codes, "S1")
        self.assertEqual(r.duration_min, 90)

    def test_la_web_ya_no_ofrece_esa_mesa(self):
        services.seat_walk_in(self.venue, self.t1, 2, now=self.ahora())
        inicio = self.ahora() + timedelta(minutes=30)
        ofertas = services.find_availability(
            self.venue, inicio, inicio + timedelta(minutes=90), 2)
        mesas = {t.code for o in ofertas for t in o.tables}
        self.assertNotIn("S1", mesas)

    def test_con_telefono_usa_la_ficha_del_cliente(self):
        r = services.seat_walk_in(self.venue, self.t1, 2, name="Pedro Soto",
                                  phone="987 555 444", now=self.ahora())
        self.assertEqual(r.guest.phone, "+51987555444")
        self.assertEqual(r.guest.name, "Pedro Soto")

    def test_se_corta_antes_de_la_siguiente_reserva(self):
        services.create_reservation(
            venue=self.venue, guest=self.guest, party_size=2, tables=[self.t1],
            starts_at=self.ahora() + timedelta(minutes=60),
            source=Reservation.Source.RECEPCION,
        )
        r = services.seat_walk_in(self.venue, self.t1, 2, now=self.ahora())
        self.assertEqual(r.duration_min, 60)
        self.assertIsNotNone(r.cut_at)

    def test_no_sienta_si_la_siguiente_reserva_esta_encima(self):
        services.create_reservation(
            venue=self.venue, guest=self.guest, party_size=2, tables=[self.t1],
            starts_at=self.ahora() + timedelta(minutes=20),
            source=Reservation.Source.RECEPCION,
        )
        with self.assertRaises(services.NoAvailability):
            services.seat_walk_in(self.venue, self.t1, 2, now=self.ahora())

    def test_no_sienta_en_una_mesa_ocupada(self):
        services.seat_walk_in(self.venue, self.t1, 2, now=self.ahora())
        with self.assertRaises(services.NoAvailability):
            services.seat_walk_in(self.venue, self.t1, 2,
                                  now=self.ahora() + timedelta(minutes=10))

    def test_no_pasa_de_la_capacidad_de_la_mesa(self):
        with self.assertRaises(services.ReservationError):
            services.seat_walk_in(self.venue, self.t1, 6, now=self.ahora())



class DetallesYMesasTests(BaseSalaTestCase):
    def test_los_detalles_se_guardan_en_la_reserva(self):
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=4,
            source=Reservation.Source.PHONE, tables=[self.t1],
            details=services.ReservationDetails(
                children=1, high_chairs=1, occasion="Cumpleanos",
                guest_notes="Torta al final", internal_notes="Cliente habitual",
            ),
        )
        reserva.refresh_from_db()
        self.assertEqual((reserva.children, reserva.high_chairs), (1, 1))
        self.assertEqual(reserva.occasion, "Cumpleanos")
        self.assertEqual(reserva.guest_notes, "Torta al final")
        self.assertEqual(reserva.internal_notes, "Cliente habitual")

    def test_mesas_elegidas_que_no_alcanzan(self):
        with self.assertRaisesMessage(services.ReservationError, "suman 4 plazas"):
            services.create_reservation(
                venue=self.venue, guest=self.guest, starts_at=self.cena(),
                party_size=6, source=Reservation.Source.PHONE, tables=[self.t1],
            )
        self.assertFalse(Reservation.objects.exists())

    def test_juntar_mesas_elegidas_para_un_grupo_grande(self):
        # S1+S2 no esta declarada para 9, pero el anfitrion puede juntarlas igual.
        t4 = Table.objects.create(area=self.area, code="S4", min_seats=1, max_seats=2)
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=9,
            source=Reservation.Source.PHONE, tables=[self.t1, self.t2, t4],
        )
        self.assertEqual(reserva.table_codes, "S1, S2, S4")
        self.assertEqual(reserva.occupancies.filter(is_primary=True).get().table, self.t1)

    def test_no_junta_una_mesa_que_no_es_combinable(self):
        Table.objects.filter(pk=self.t2.pk).update(is_combinable=False)
        self.t2.refresh_from_db()
        with self.assertRaisesMessage(services.ReservationError, "S2 no se puede juntar"):
            services.create_reservation(
                venue=self.venue, guest=self.guest, starts_at=self.cena(),
                party_size=6, source=Reservation.Source.PHONE, tables=[self.t1, self.t2],
            )
        self.assertFalse(Reservation.objects.exists())

    def test_mesa_elegida_ocupada_avisa_cual(self):
        services.block_table(self.t2, self.cena(hora=19), self.cena(hora=23),
                             reason="Evento")
        with self.assertRaisesMessage(services.NoAvailability, "S2 ya está ocupada"):
            services.create_reservation(
                venue=self.venue, guest=self.guest, starts_at=self.cena(),
                party_size=6, source=Reservation.Source.PHONE, tables=[self.t1, self.t2],
            )

    def test_mover_conserva_la_mesa_si_sigue_libre(self):
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        movida = services.move(reserva, self.cena(hora=21))
        self.assertEqual(timezone.localtime(movida.starts_at, LIMA).hour, 21)
        self.assertEqual(movida.table_codes, "S1")
        self.assertTrue(movida.events.filter(comment__startswith="Movida").exists())

    def test_mover_busca_otra_mesa_si_la_suya_esta_ocupada(self):
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        # La reserva ocupa S1 hasta las 21:30; desde ahi S1 esta bloqueada.
        services.block_table(self.t1, self.cena(hora=21, minuto=30), self.cena(hora=23),
                             reason="Toldo roto")
        movida = services.move(reserva, self.cena(hora=21, minuto=30))
        self.assertEqual(movida.table_codes, "S2")

    def test_mover_fuera_de_turno_falla(self):
        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.PHONE, tables=[self.t1],
        )
        with self.assertRaises(services.NoAvailability):
            services.move(reserva, self.cena(hora=17))


class AdminReservaFormTests(BaseSalaTestCase):
    def setUp(self):
        from django.contrib.auth.models import Permission, User

        self.staff = User.objects.create_user("host", password=secrets.token_urlsafe(),
                                              is_staff=True)
        self.staff.user_permissions.add(*Permission.objects.filter(
            codename__in=["view_reservation", "add_reservation", "change_reservation"]))
        self.client.force_login(self.staff)

    def formulario(self, inicio):
        from apps.reservations.admin import ReservationForm

        local = timezone.localtime(inicio)
        return ReservationForm(data={
            "venue": self.venue.pk, "guest": self.guest.pk,
            "starts_at_0": f"{local:%Y-%m-%d}", "starts_at_1": f"{local:%H:%M}",
            "party_size": 2, "source": Reservation.Source.RECEPCION,
        })

    def test_no_se_reserva_en_un_dia_que_ya_paso(self):
        form = self.formulario(self.cena(dias_desde_hoy=-1))
        self.assertFalse(form.is_valid())
        self.assertIn("No se puede reservar en un día que ya pasó", form.errors[STARTS_AT][0])

    def test_tampoco_se_mueve_una_reserva_a_un_dia_pasado(self):
        from apps.reservations.admin import ReservationForm

        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.RECEPCION,
        )
        ayer = timezone.localtime(self.cena(dias_desde_hoy=-1))
        form = ReservationForm(instance=reserva, data={
            "venue": self.venue.pk, "guest": self.guest.pk,
            "starts_at_0": f"{ayer:%Y-%m-%d}", "starts_at_1": f"{ayer:%H:%M}",
            "party_size": 2, "source": Reservation.Source.RECEPCION,
            "status": reserva.status, "deposit_status": reserva.deposit_status,
            "children": 0, "high_chairs": 0,
        })
        self.assertFalse(form.is_valid())
        self.assertIn("ya pasó", form.errors[STARTS_AT][0])

    def test_tampoco_una_hora_de_hoy_que_ya_paso(self):
        from unittest import mock

        a_las_21 = self.cena(dias_desde_hoy=0, hora=21)
        with mock.patch("django.utils.timezone.now", return_value=a_las_21):
            form = self.formulario(self.cena(dias_desde_hoy=0, hora=20))
            self.assertFalse(form.is_valid())
            self.assertIn("Esa hora ya pasó: son las 21:00", form.errors[STARTS_AT][0])

    def test_una_hora_de_hoy_que_aun_no_llega_si(self):
        from unittest import mock

        a_las_19 = self.cena(dias_desde_hoy=0, hora=19)
        with mock.patch("django.utils.timezone.now", return_value=a_las_19):
            form = self.formulario(self.cena(dias_desde_hoy=0, hora=20, minuto=30))
            form.is_valid()
            self.assertNotIn(STARTS_AT, form.errors)

    def test_horas_cada_media_hora_de_6_a_22(self):
        from apps.reservations.admin import ReservationForm

        horas = [h for h, _ in ReservationForm().fields[STARTS_AT].widget.widgets[1].choices]
        self.assertEqual((horas[0], horas[1], horas[-1], len(horas)), ("06:00", "06:30", "22:00", 33))

    def test_una_hora_fuera_de_la_lista_no_se_pierde_al_editar(self):
        from apps.reservations.admin import ReservationForm

        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(hora=20, minuto=15),
            party_size=2, source=Reservation.Source.RECEPCION,
        )
        html = str(ReservationForm(instance=reserva)[STARTS_AT])
        self.assertIn('<option value="20:15" selected>', html)
        self.assertIn('<option value="20:30">', html)

    def test_en_una_reserva_nueva_hay_que_elegir_la_hora(self):
        from apps.reservations.admin import ReservationForm

        html = str(ReservationForm()[STARTS_AT])
        self.assertIn('<option value="">Elige la hora</option>', html)
        self.assertNotIn("selected", html)
        form = self.formulario(self.cena())
        form.data = {**form.data, "starts_at_1": ""}
        self.assertFalse(form.is_valid())
        self.assertIn(STARTS_AT, form.errors)

    def test_hora_fuera_de_turno_es_error_del_campo(self):
        form = self.formulario(self.cena(hora=17))
        self.assertFalse(form.is_valid())
        self.assertIn("No hay turno abierto", form.errors[STARTS_AT][0])

    def test_sin_mesa_libre_es_error_del_campo(self):
        for mesa in (self.t1, self.t2):
            services.block_table(mesa, self.cena(hora=19), self.cena(hora=23),
                                 reason="Evento")
        form = self.formulario(self.cena())
        self.assertFalse(form.is_valid())
        self.assertIn("No queda mesa libre para 2 personas", form.errors[STARTS_AT][0])

    def test_una_mesa_elegida_en_el_alta(self):
        form = self.formulario(self.cena())
        form.data = {**form.data, "table": f"t:{self.t2.pk}"}
        form.is_valid()
        self.assertNotIn("table", form.errors)
        self.assertEqual(form.cleaned_data["chosen_tables"], [self.t2])

    def test_sin_juntar_se_ignora_la_lista_de_varias(self):
        form = self.formulario(self.cena())
        form.data = {**form.data, "tables": [str(self.t1.pk), str(self.t2.pk)]}
        form.is_valid()
        self.assertIsNone(form.cleaned_data["chosen_tables"])

    def test_juntar_mesas_en_el_alta(self):
        form = self.formulario(self.cena())
        form.data = {**form.data, "party_size": 7, "combine": "on",
                     "tables": [str(self.t1.pk), str(self.t2.pk)]}
        form.is_valid()
        self.assertNotIn("tables", form.errors)
        self.assertNotIn(STARTS_AT, form.errors)
        self.assertEqual(set(form.cleaned_data["chosen_tables"]), {self.t1, self.t2})

    def test_mesas_juntadas_que_no_alcanzan_es_error_del_campo(self):
        form = self.formulario(self.cena())
        form.data = {**form.data, "party_size": 10, "combine": "on",
                     "tables": [str(self.t1.pk), str(self.t2.pk)]}
        self.assertFalse(form.is_valid())
        self.assertIn("suman 8 plazas", form.errors["tables"][0])

    def test_mesa_unica_ocupada_es_error_de_su_campo(self):
        services.block_table(self.t1, self.cena(hora=19), self.cena(hora=23),
                             reason="Evento")
        form = self.formulario(self.cena())
        form.data = {**form.data, "table": f"t:{self.t1.pk}"}
        self.assertFalse(form.is_valid())
        self.assertIn("S1 ya está ocupada", form.errors["table"][0])

    def test_la_combinacion_declarada_sale_en_la_lista(self):
        from apps.reservations.admin import ReservationForm

        html = str(ReservationForm()["table"])
        self.assertIn('label="Mesas que ya se juntan"', html)
        self.assertIn(f'value="c:{self.combo.pk}"', html)
        self.assertIn("S1+S2 · juntas (5-8)", html)

    def test_elegir_la_combinacion_declarada(self):
        form = self.formulario(self.cena())
        form.data = {**form.data, "party_size": 6, "table": f"c:{self.combo.pk}"}
        form.is_valid()
        self.assertNotIn("table", form.errors)
        self.assertEqual(set(form.cleaned_data["chosen_tables"]), {self.t1, self.t2})

    def test_combinacion_declarada_respeta_su_aforo(self):
        form = self.formulario(self.cena())
        form.data = {**form.data, "party_size": 9, "table": f"c:{self.combo.pk}"}
        self.assertFalse(form.is_valid())
        self.assertIn("S1+S2 es para 8 personas como máximo", form.errors["table"][0])

    def test_combinacion_inactiva_no_sale(self):
        from apps.reservations.admin import ReservationForm

        TableCombination.objects.filter(pk=self.combo.pk).update(is_active=False)
        self.assertNotIn(f"c:{self.combo.pk}", str(ReservationForm()["table"]))

    def test_las_opciones_de_mesa_llevan_su_zona(self):
        from apps.reservations.admin import ReservationForm

        html = str(ReservationForm()["table"])
        self.assertIn(f'data-area="{self.area.pk}"', html)
        self.assertIn("Que la elija el sistema", html)

    def test_alta_desde_el_admin_junta_las_mesas_marcadas(self):
        self.staff.is_superuser = True
        self.staff.save()
        local = timezone.localtime(self.cena())
        respuesta = self.client.post(reverse("admin:reservations_reservation_add"), {
            "venue": self.venue.pk, "guest": self.guest.pk,
            "starts_at_0": f"{local:%Y-%m-%d}", "starts_at_1": f"{local:%H:%M}",
            "party_size": 8, "children": 0, "high_chairs": 0,
            "source": Reservation.Source.RECEPCION, "combine": "on",
            "tables": [str(self.t1.pk), str(self.t2.pk)],
        })
        self.assertEqual(respuesta.status_code, 302, getattr(respuesta, "context", None)
                         and respuesta.context.get("errors"))
        reserva = Reservation.objects.get()
        self.assertEqual(reserva.table_codes, "S1, S2")

    def test_campos_de_solo_lectura_en_alta_y_en_ficha(self):
        from django.contrib import admin
        from django.urls import reverse

        reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.RECEPCION,
        )
        ma = admin.site._registry[Reservation]
        self.assertEqual(ma.get_readonly_fields(None), [])
        self.assertIn("status", ma.get_readonly_fields(None, reserva))

        alta = self.client.get(reverse("admin:reservations_reservation_add"))
        self.assertEqual(alta.status_code, 200)
        ficha = self.client.get(reverse("admin:reservations_reservation_change",
                                        args=[reserva.pk]))
        self.assertEqual(ficha.status_code, 200)
        self.assertContains(ficha, reserva.code)


class SeedDemoDayTests(BaseSalaTestCase):
    def correr(self, *args):
        from io import StringIO

        from django.core.management import call_command

        salida = StringIO()
        call_command("seed_demo_day", *args, stdout=salida)
        return salida.getvalue()

    def test_crea_rehace_y_borra_las_de_ejemplo(self):
        from apps.reservations.management.commands.seed_demo_day import DEMO_TAG

        demo = Reservation.objects.filter(internal_notes__startswith=DEMO_TAG)
        salida = self.correr()
        self.assertIn("reservas de ejemplo creadas", salida)
        creadas = demo.count()
        self.assertGreater(creadas, 0)
        # Ocasion y notas llegan por ReservationDetails.
        self.assertTrue(demo.exclude(occasion="").exists())

        salida = self.correr("--venue", self.venue.name)
        self.assertIn(f"{creadas} reservas de ejemplo anteriores borradas", salida)

        salida = self.correr("--clear")
        self.assertIn("reservas de ejemplo borradas", salida)
        self.assertFalse(demo.exists())

    def test_sin_sede_avisa(self):
        from django.core.management.base import CommandError

        Venue.objects.update(is_active=False)
        with self.assertRaises(CommandError):
            self.correr()


class BorrarClienteTests(BaseSalaTestCase):
    """El superadmin borra al cliente con todo su historial; el resto del equipo no."""

    def setUp(self):
        from django.contrib.auth.models import User

        self.jefe = User.objects.create_superuser("jefe", password=secrets.token_urlsafe())
        self.reserva = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(), party_size=2,
            source=Reservation.Source.RECEPCION,
        )
        WaitlistEntry.objects.create(venue=self.venue, guest=self.guest, date=date.today(),
                                     desired_time=time(20, 0), party_size=2)
        self.url = reverse("admin:guests_guest_delete", args=[self.guest.pk])

    def test_el_superadmin_borra_al_cliente_y_su_historial(self):
        self.client.force_login(self.jefe)
        pagina = self.client.get(self.url)
        self.assertEqual(pagina.status_code, 200)
        self.assertFalse(pagina.context["protected"])

        self.client.post(self.url, {"post": "yes"})
        self.assertFalse(Guest.objects.filter(pk=self.guest.pk).exists())
        self.assertFalse(Reservation.objects.filter(pk=self.reserva.pk).exists())
        self.assertFalse(WaitlistEntry.objects.exists())
        self.assertFalse(TableOccupancy.objects.filter(reservation_id=self.reserva.pk).exists())

    def test_el_equipo_no_puede_borrar_clientes(self):
        from django.contrib.auth.models import Permission, User

        staff = User.objects.create_user("host", password=secrets.token_urlsafe(),
                                         is_staff=True)
        staff.user_permissions.add(*Permission.objects.filter(
            codename__in=["view_guest", "delete_guest"]))
        self.client.force_login(staff)
        self.assertEqual(self.client.post(self.url, {"post": "yes"}).status_code, 403)
        self.assertTrue(Guest.objects.filter(pk=self.guest.pk).exists())


class ExcelTests(BaseSalaTestCase):
    """Descarga en Excel por rango de fechas, solo para el superusuario."""

    RESERVAS = "admin:reservations_reservation_export_excel"
    CLIENTES = "admin:guests_guest_export_excel"

    def setUp(self):
        from django.contrib.auth.models import User

        self.jefe = User.objects.create_superuser("jefe", password=secrets.token_urlsafe())
        self.manana = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(hora=21), party_size=3,
            source=Reservation.Source.RECEPCION,
        )
        self.pasado = services.create_reservation(
            venue=self.venue, guest=self.guest, starts_at=self.cena(2), party_size=2,
            source=Reservation.Source.RECEPCION,
        )
        services.cancel(self.pasado, reason="Cambio de planes")
        self.client.force_login(self.jefe)

    def excel(self, url, **params):
        from io import BytesIO

        from openpyxl import load_workbook

        r = self.client.get(reverse(url), params)
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment;", r["Content-Disposition"])
        return load_workbook(BytesIO(r.content)), r["Content-Disposition"]

    def test_sin_rango_muestra_la_pantalla_para_elegirlo(self):
        r = self.client.get(reverse(self.RESERVAS))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'name="desde"')
        self.assertContains(r, "Semana pasada")
        self.assertNotIn("Content-Disposition", r)

    def test_rango_al_reves_avisa(self):
        r = self.client.get(reverse(self.RESERVAS), {"desde": "2026-09-30", "hasta": "2026-09-01"})
        self.assertContains(r, "no puede ser posterior")

    def test_reservas_de_un_dia_y_resumen(self):
        dia = self.manana.service_date.isoformat()
        wb, cabecera = self.excel(self.RESERVAS, desde=dia, hasta=dia)
        self.assertIn(f"reservas_{dia}_a_{dia}.xlsx", cabecera)
        filas = list(wb["Reservas"].iter_rows(values_only=True))
        self.assertEqual(filas[0][:5], ("Fecha", "Hora", "Hasta", "Código", "Cliente"))
        self.assertEqual([f[3] for f in filas[1:]], [self.manana.code])
        self.assertEqual(filas[1][1].strftime("%H:%M"), "21:00")
        resumen = list(wb["Resumen por día"].iter_rows(values_only=True))
        self.assertEqual(resumen[1][2:], (1, 3, 0, 0))

    def test_rango_de_varios_dias_la_cancelada_no_suma_personas(self):
        wb, _ = self.excel(self.RESERVAS, desde=self.manana.service_date.isoformat(),
                           hasta=self.pasado.service_date.isoformat())
        self.assertEqual(wb["Reservas"].max_row, 3)
        resumen = {f[0].date(): f[2:] for f in
                   wb["Resumen por día"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(resumen[self.pasado.service_date], (1, 0, 1, 0))

    def test_clientes_que_reservaron_o_nuevos(self):
        otro = Guest.objects.create(phone="+51911111111", name="Beto Sin Reservas")
        dia = self.manana.service_date.isoformat()
        wb, _ = self.excel(self.CLIENTES, desde=dia, hasta=dia, modo="reservaron")
        nombres = [f[0] for f in wb["Clientes"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(nombres, ["Ana Rojas"])

        hoy = timezone.localdate().isoformat()
        wb, _ = self.excel(self.CLIENTES, desde=hoy, hasta=hoy, modo="nuevos")
        nombres = [f[0] for f in wb["Clientes"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(nombres), ["Ana Rojas", otro.name])

    def test_solo_el_superadmin_ve_el_boton_y_descarga(self):
        from django.contrib.auth.models import Permission, User

        self.assertContains(self.client.get(reverse("admin:guests_guest_changelist")),
                            "Descargar Excel")
        staff = User.objects.create_user("host", password=secrets.token_urlsafe(),
                                         is_staff=True)
        staff.user_permissions.add(*Permission.objects.filter(
            codename__in=["view_guest", "view_reservation"]))
        self.client.force_login(staff)
        for url in (self.CLIENTES, self.RESERVAS):
            self.assertEqual(self.client.get(reverse(url)).status_code, 403)
        for lista in ("admin:guests_guest_changelist", "admin:reservations_reservation_changelist"):
            self.assertNotContains(self.client.get(reverse(lista)), "Descargar Excel")
