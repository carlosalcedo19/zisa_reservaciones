"""Descargas en Excel desde las listas del panel, por rango de fechas."""

from datetime import timedelta
from io import BytesIO

from django import forms
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from unfold.decorators import action

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="18181B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FORMATS = {"date": "dd/mm/yyyy", "time": "hh:mm", "datetime": "dd/mm/yyyy hh:mm"}


def add_sheet(wb, title, columns, rows):
    """columns: [(titulo, ancho, formato o None)]; rows: listas de valores en ese orden."""
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, (_, width, fmt) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = FORMATS[fmt]
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def workbook_response(wb, filename):
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX)
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def range_presets(today):
    week = today - timedelta(days=today.weekday())
    month = today.replace(day=1)
    prev_month_end = month - timedelta(days=1)
    return [
        ("Hoy", today, today),
        ("Ayer", today - timedelta(days=1), today - timedelta(days=1)),
        ("Esta semana", week, week + timedelta(days=6)),
        ("Semana pasada", week - timedelta(days=7), week - timedelta(days=1)),
        ("Este mes", month, (month + timedelta(days=32)).replace(day=1) - timedelta(days=1)),
        ("Mes pasado", prev_month_end.replace(day=1), prev_month_end),
    ]


class RangeForm(forms.Form):
    desde = forms.DateField(label="Desde")
    hasta = forms.DateField(label="Hasta")

    def clean(self):
        data = super().clean()
        if data.get("desde") and data.get("hasta") and data["desde"] > data["hasta"]:
            raise forms.ValidationError("La fecha «Desde» no puede ser posterior a «Hasta».")
        return data


class ExcelExportMixin:
    """
    Boton "Descargar Excel" en la lista, solo para el superusuario. Abre una
    pantalla para elegir el rango de fechas; cada admin define que significa
    ese rango (excel_filter) y que columnas lleva el archivo (excel_workbook).
    """

    excel_name = "datos"
    excel_form = RangeForm
    excel_help = ""

    # Datos personales de todos los clientes de golpe: solo el superusuario.
    def has_export_permission(self, request):
        return request.user.is_superuser

    def excel_filter(self, queryset, data):
        raise NotImplementedError

    def excel_workbook(self, request, queryset):
        raise NotImplementedError

    @action(description="Descargar Excel", icon="download", url_path="excel",
            permissions=["export"])
    def export_excel(self, request):
        today = timezone.localdate()
        month = today.replace(day=1)
        form = self.excel_form(request.GET or None, initial={"desde": month, "hasta": today})
        if form.is_bound and form.is_valid():
            data = form.cleaned_data
            queryset = self.excel_filter(self.get_queryset(request), data)
            name = f"{self.excel_name}_{data['desde']:%Y-%m-%d}_a_{data['hasta']:%Y-%m-%d}"
            return workbook_response(self.excel_workbook(request, queryset), name)

        opts = self.model._meta
        return render(request, "admin/excel_export.html", {
            **self.admin_site.each_context(request),
            "title": f"Descargar {opts.verbose_name_plural} en Excel",
            "opts": opts,
            "form": form,
            "help": self.excel_help,
            "presets": [(label, f"{a:%Y-%m-%d}", f"{b:%Y-%m-%d}")
                        for label, a, b in range_presets(today)],
            "back_url": reverse(f"admin:{opts.app_label}_{opts.model_name}_changelist"),
        })
